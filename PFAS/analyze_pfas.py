"""Reproduce PFAS fingerprint-array LDA, response bars and confusion matrices."""
import argparse
from pathlib import Path
import shutil
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.patches import Ellipse, Patch
from matplotlib.backends.backend_pdf import PdfPages
from scipy.stats import chi2
from sklearn.discriminant_analysis import LinearDiscriminantAnalysis
from sklearn.preprocessing import StandardScaler
from sklearn.metrics import confusion_matrix
from openpyxl import load_workbook, Workbook

CHANNELS = ["WT", "LA", "LB", "LC", "LD", "LE"]
CHANNEL_COLORS = ["#4FB4D0", "#E64E3C", "#009C8A", "#40598A", "#F19C83", "#8490B2"]
CLASS_COLORS = ["#E64E3C", "#4FB4D0", "#009C8A", "#40598A", "#F19C83"]
GROUPS = [
    ("Five_PFAS", "Five PFAS", 0, 1, "PFAS"),
    ("PFOA_PFOS_mixture", "PFOA/PFOS mixture", 2, 3, "PFOA:PFOS ratio"),
    ("PFOA_concentration", "PFOA concentration", 4, 5, "Concentration (μM)"),
    ("PFOS_concentration", "PFOS concentration", 6, 7, "Concentration (μM)"),
]
plt.rcParams.update(dict(font_family="DejaVu Sans") if False else {
    "font.family": "DejaVu Sans", "font.size": 12, "axes.titlesize": 16,
    "axes.labelsize": 14, "axes.linewidth": 1.2,
    "pdf.fonttype": 42, "svg.fonttype": "none",
    "savefig.facecolor": "white", "figure.facecolor": "white",
})

def clean_label(value):
    return str(value).replace("：", ":").replace(" μM", "")

def save_figure(fig, folder, slug, kind, document):
    for ext in ("png", "pdf", "svg"):
        fig.savefig(folder / f"{slug}_{kind}.{ext}", dpi=300, facecolor="white")
    document.savefig(fig)
    plt.close(fig)

def setup_ax(fig, rect):
    ax = fig.add_axes(rect)
    ax.spines["right"].set_visible(False)
    ax.spines["top"].set_visible(False)
    return ax

def read_sheet(sheet):
    return [list(row) for row in sheet.iter_rows(values_only=True) if any(v is not None for v in row)]

def lda_plot(slug, title, classes, train_y, test_y, ztrain, ztest, ratios, dest, pdf):
    fig = plt.figure(figsize=(9.6, 5.8))
    ax = setup_ax(fig, [0.10, 0.14, 0.60, 0.74])
    shapes = []
    for k, lab in enumerate(classes):
        pts = ztrain[train_y == k]
        mean = pts.mean(axis=0)
        vals, vecs = np.linalg.eigh(np.cov(pts, rowvar=False))
        radii = np.sqrt(np.maximum(vals, 1e-8) * chi2.ppf(.95, 2))
        angle = np.rad2deg(np.arctan2(vecs[1, 1], vecs[0, 1]))
        ax.add_patch(Ellipse(mean, 2*radii[1], 2*radii[0], angle=angle,
                             facecolor=CLASS_COLORS[k], alpha=.12, edgecolor="none"))
        theta = np.linspace(0, 2*np.pi, 180)
        shapes.append(mean + (vecs @ (radii[:, None] * np.array([np.cos(theta), np.sin(theta)]))).T)
        ax.scatter(*pts.T, s=48, facecolor=CLASS_COLORS[k],
                   edgecolor="#2f2f2f", linewidth=.5, zorder=3)
        test = ztest[test_y == k]
        ax.scatter(*test.T, s=66, facecolor="white",
                   edgecolor=CLASS_COLORS[k], linewidth=1.6, zorder=4)
    whole = np.vstack([ztrain, ztest] + shapes)
    lo, hi = whole.min(axis=0), whole.max(axis=0)
    pad = np.maximum((hi-lo)*.12, .8)
    ax.set_xlim(lo[0]-pad[0], hi[0]+pad[0])
    ax.set_ylim(lo[1]-pad[1], hi[1]+pad[1])
    ax.set(xlabel=f"LD1 ({ratios[0]*100:.1f}%)", ylabel=f"LD2 ({ratios[1]*100:.1f}%)")
    ax.set_title(title + " discrimination", fontweight="bold", pad=10)
    category_handles = [
        Line2D([], [], marker="o", linestyle="none", markersize=7.5,
               markerfacecolor=CLASS_COLORS[k], markeredgecolor="#333333",
               label=lab + (" μM" if "concentration" in slug else ""))
        for k, lab in enumerate(classes)
    ]
    fig.legend(handles=category_handles, frameon=False, fontsize=10.5,
               loc="upper right", bbox_to_anchor=(.975, .90))
    split_handles = [
        Line2D([], [], marker="o", linestyle="none", markersize=7,
               markerfacecolor="#666666", markeredgecolor="#333333", label="Training (n=6)"),
        Line2D([], [], marker="o", linestyle="none", markersize=8,
               markerfacecolor="white", markeredgecolor="#666666",
               markeredgewidth=1.5, label="Prediction (n=4)"),
    ]
    fig.legend(handles=split_handles, frameon=False, fontsize=10.5,
               loc="upper right", bbox_to_anchor=(.985, .48))
    save_figure(fig, dest, slug, "LDA", pdf)

def bars_plot(slug, title, labels, means, errs, xlabel, dest, pdf):
    fig = plt.figure(figsize=(10.6, 5.8))
    ax = setup_ax(fig, [.09, .16, .62, .70])
    x = np.arange(len(labels))
    for j, channel in enumerate(CHANNELS):
        positions = x + (j-2.5)*.13
        ax.bar(positions, means[:, j], width=.13*.95,
               color=CHANNEL_COLORS[j], edgecolor="#444444", linewidth=.45)
        ax.errorbar(positions, means[:, j], yerr=errs[:, j], fmt="none",
                    ecolor="#4b4b4b", elinewidth=.9, capsize=2.6, capthick=.9)
    low, high = np.min(means-errs), np.max(means+errs)
    ax.set_ylim(low - max(.08*(high-low), .015),
                high + max(.18*(high-low), .03, .06 if high <= 0 else 0))
    ax.axhline(0, linewidth=.75, color="#333333")
    ax.set_xticks(x, labels)
    ax.set(xlabel=xlabel, ylabel="Normalized fluorescence response")
    ax.set_title(title + " response profile", fontweight="bold", pad=12)
    handles = [Patch(facecolor=CHANNEL_COLORS[j], edgecolor="#444444",
                     linewidth=.45, label=ch) for j, ch in enumerate(CHANNELS)]
    fig.legend(handles=handles, loc="upper right", bbox_to_anchor=(.975, .84),
               ncol=1, frameon=False, fontsize=10.5)
    save_figure(fig, dest, slug, "Bar", pdf)

def matrix_plot(slug, classes, cm, dest, pdf):
    count = int(np.trace(cm))
    fig = plt.figure(figsize=(6.6, 5.8))
    ax = setup_ax(fig, [.15, .14, .72, .72])
    ax.imshow(cm, cmap="Blues", vmin=0, vmax=4)
    for (i, j), num in np.ndenumerate(cm):
        ax.text(j, i, str(num), ha="center", va="center", fontweight="bold",
                color="white" if num >= 3 else "#213046")
    rot = 28 if slug == "Five_PFAS" else 0
    ax.set_xticks(range(len(classes)), classes, rotation=rot,
                  ha="right" if rot else "center")
    ax.set_yticks(range(len(classes)), classes)
    ax.set(xlabel="Predicted class", ylabel="True class")
    ax.set_title(f"Blind-test confusion matrix ({count}/20, {count/20:.0%})",
                 fontweight="bold", pad=10)
    save_figure(fig, dest, slug, "Confusion", pdf)

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=Path(__file__).parent/"20261017-ALL LDA and Bar chart.xlsx")
    parser.add_argument("--output", type=Path, default=Path(__file__).parent/"results")
    parser.add_argument("--seed", type=int, default=20261017)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    folder = args.output/"individual_figures"
    folder.mkdir(exist_ok=True)
    source = load_workbook(args.input, read_only=True, data_only=True)
    if len(source.worksheets) < 8:
        raise ValueError("Expected 8 worksheets: four LDA and four bar tables.")
    summary = Workbook()
    header = ["Dataset","Classes","Train","Test","Correct","Accuracy","LD1 (%)","LD2 (%)","Seed"]
    summary.active.title = "Summary"
    summary.active.append(header)
    with PdfPages(args.output/"PFAS_all_figures.pdf") as pdf:
        for slug, title, di, bi, xlabel in GROUPS:
            records = read_sheet(source.worksheets[di])
            labels = [clean_label(row[0]) for row in records]
            classes = list(dict.fromkeys(labels))
            y = np.array([classes.index(v) for v in labels])
            X = np.array([row[1:7] for row in records], dtype=float)
            rng = np.random.default_rng(args.seed)
            train, test = [], []
            for k in range(len(classes)):
                indices = rng.permutation(np.flatnonzero(y == k))
                train.extend(indices[:6])
                test.extend(indices[6:])
            train, test = np.array(train), np.array(test)
            scaler = StandardScaler().fit(X[train])
            Xtrain, Xtest = scaler.transform(X[train]), scaler.transform(X[test])
            lda = LinearDiscriminantAnalysis(n_components=2).fit(Xtrain, y[train])
            ztrain, ztest = lda.transform(Xtrain), lda.transform(Xtest)
            predictions = lda.predict(Xtest)
            cm = confusion_matrix(y[test], predictions, labels=range(len(classes)))
            ratios = lda.explained_variance_ratio_
            summary.active.append([title, len(classes), len(train), len(test),
                                   int(np.trace(cm)), float(np.trace(cm)/len(test)),
                                   float(ratios[0]*100), float(ratios[1]*100), args.seed])
            table = read_sheet(source.worksheets[bi])[1:]
            barlabels = [clean_label(row[0]) for row in table]
            means = np.array([[row[1+2*j] for j in range(6)] for row in table], dtype=float)
            errors = np.array([[row[2+2*j] for j in range(6)] for row in table], dtype=float)
            lda_plot(slug, title, classes, y[train], y[test], ztrain, ztest, ratios, folder, pdf)
            bars_plot(slug, title, barlabels, means, errors, xlabel, folder, pdf)
            matrix_plot(slug, classes, cm, folder, pdf)
            ws = summary.create_sheet(slug[:31])
            ws.append(["Original row","Class",*CHANNELS,"Split","LD1","LD2","Prediction"])
            z = np.zeros((len(y), 2))
            z[train], z[test] = ztrain, ztest
            predmap = dict(zip(test, predictions))
            trainset = set(train)
            for i, row in enumerate(X):
                ws.append([i+1, labels[i], *map(float, row),
                           "Training" if i in trainset else "Prediction",
                           float(z[i, 0]), float(z[i, 1]),
                           classes[predmap[i]] if i in predmap else ""])
    source.close()
    summary.save(args.output/"PFAS_LDA_results.xlsx")
    shutil.make_archive(str(args.output/"individual_figures"), "zip", folder)
    print("Wrote figures and data to", args.output)

if __name__ == "__main__":
    main()
